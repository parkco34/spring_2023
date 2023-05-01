#!/usr/bin/env python
import os
import re
import openpyxl
import csv
import docx

def search_in_txt(file_path, search_word):
    with open(file_path, 'r') as file:
        contents = file.read()
        return search_word in contents

def search_in_xlsx(file_path, search_word):
    workbook = openpyxl.load_workbook(file_path)
    for sheet in workbook:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and search_word in str(cell.value):
                    return True
    return False

def search_in_csv(file_path, search_word):
    with open(file_path, newline='') as csvfile:
        reader = csv.reader(csvfile)
        for row in reader:
            for item in row:
                if search_word in item:
                    return True
    return False

def search_in_docx(file_path, search_word):
    doc = docx.Document(file_path)
    for paragraph in doc.paragraphs:
        if search_word in paragraph.text:
            return True
    return False

def search_word_in_file(file_path, search_word):
    _, ext = os.path.splitext(file_path)
    if ext.lower() in {'.txt', '.cpp', '.c', '.java', '.py', '.js', '.dat', '.cc'}:
        return search_in_txt(file_path, search_word)
    elif ext.lower() == '.xlsx':
        return search_in_xlsx(file_path, search_word)
    elif ext.lower() == '.csv':
        return search_in_csv(file_path, search_word)
    elif ext.lower() == '.docx':
        return search_in_docx(file_path, search_word)
    return False

def search_in_directory(search_directory, search_word):
    for root, _, files in os.walk(search_directory):
        for file in files:
            file_path = os.path.join(root, file)
            if search_word_in_file(file_path, search_word):
                return file_path
    return None

def main():
    search_directory = input("Enter the directory path: ")
    search_word = input("Enter the word you want to search for: ")

    found_file = search_in_directory(search_directory, search_word)

    if found_file:
        print(f"Found {search_word} in the file: {found_file}")
    else:
        print(f"{search_word} not found in any file.")

if __name__ == "__main__":
    main()

